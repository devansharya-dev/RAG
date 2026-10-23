
"""Microsoft Word DOCX document loader."""

from pathlib import Path

from docx import Document

from app.ingestion.loaders.base import BaseLoader, LoadedDocument


class DOCXLoader(BaseLoader):
    """Extract paragraphs and tables from DOCX files."""

    supported_extensions = {".docx"}

    def load(self, file_path: Path) -> list[LoadedDocument]:
        file_path = Path(file_path)

        if not file_path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        if not self.supports(file_path):
            raise ValueError(
                f"Unsupported Word document extension: {file_path.suffix}"
            )

        try:
            document = Document(file_path)
        except Exception as exc:
            raise ValueError(
                f"Could not read DOCX '{file_path.name}': {exc}"
            ) from exc

        sections: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(value for value in cells if value)
                if row_text:
                    sections.append(row_text)

        content = "\n".join(sections).strip()

        if not content:
            raise ValueError(
                f"No extractable text found in DOCX: {file_path.name}"
            )

        return [
            LoadedDocument(
                content=content,
                metadata={
                    "filename": file_path.name,
                    "path": str(file_path.resolve()),
                    "file_type": ".docx",
                    "file_size": file_path.stat().st_size,
                },
            )
        ]
